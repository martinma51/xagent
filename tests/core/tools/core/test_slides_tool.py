"""Tests for the slide-deck orchestrator (slides_tool)."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from xagent.core.tools.core.slides_tool import (
    _fill_html_slots,
    _validate_slot_values,
    create_deck_from_template,
    export_deck_to_pptx,
    get_slide_template,
    list_slide_templates,
    render_deck,
)


# ---------------------------------------------------------------------------
# discovery
# ---------------------------------------------------------------------------


def test_list_templates_includes_builtin_midnight_tech() -> None:
    result = list_slide_templates()
    assert result["success"] is True
    ids = {t["id"] for t in result["templates"]}
    assert "midnight_tech" in ids
    mt = next(t for t in result["templates"] if t["id"] == "midnight_tech")
    assert mt["page_count"] == 5
    # five thumbnails are bundled with the template
    assert len(mt["thumbnails"]) == 5


def test_list_templates_filters_by_category() -> None:
    tech = list_slide_templates(category="Tech")
    assert all(t["category"].lower() == "tech" for t in tech["templates"])
    none = list_slide_templates(category="Made-Up-Category")
    assert none["templates"] == []


def test_get_template_returns_slot_schema() -> None:
    meta = get_slide_template("midnight_tech")
    assert meta["success"] is True
    pages = meta["template"]["pages"]
    assert len(pages) == 5
    cover = pages[0]
    assert cover["layout"] == "cover"
    assert "title" in cover["slots"]
    assert cover["slots"]["title"]["required"] is True


def test_get_template_unknown_id() -> None:
    res = get_slide_template("does_not_exist")
    assert res["success"] is False
    assert "unknown" in res["error"]


# ---------------------------------------------------------------------------
# slot filling
# ---------------------------------------------------------------------------


def test_fill_html_slots_replaces_only_marked_elements() -> None:
    html = (
        "<html><body>"
        '<div class="title" data-slot="title">DEFAULT</div>'
        '<div class="other">untouched</div>'
        "</body></html>"
    )
    out = _fill_html_slots(html, {"title": "New Title"})
    assert "New Title" in out
    assert "untouched" in out
    assert "DEFAULT" not in out


def test_fill_html_slots_keeps_default_when_no_value() -> None:
    html = '<p data-slot="x">keep me</p>'
    out = _fill_html_slots(html, {"y": "ignored"})
    assert "keep me" in out
    assert "ignored" not in out


def test_validate_slot_values_flags_missing_required() -> None:
    meta = {
        "pages": [
            {"idx": 0, "slots": {"title": {"required": True}, "subtitle": {"required": False}}},
            {"idx": 1, "slots": {"heading": {"required": True}}},
        ]
    }
    errors = _validate_slot_values(meta, {0: {"subtitle": "x"}, 1: {"heading": "y"}})
    assert errors == ["page 0 missing required slot 'title'"]


# ---------------------------------------------------------------------------
# end-to-end
# ---------------------------------------------------------------------------


def test_render_deck_writes_filled_html(tmp_path: Path) -> None:
    slot_values = {
        0: {"title": "Hello World", "subtitle": "An end-to-end test", "author": "QA"},
        1: {"heading": "Today's Plan", "item_1": "Alpha", "item_2": "Beta", "item_3": "Gamma"},
        2: {"title": "Background", "number": "01"},
        3: {
            "heading": "What's new",
            "feat_1_title": "Speed",
            "feat_1_body": "Faster",
            "feat_2_title": "Safety",
            "feat_2_body": "Safer",
            "feat_3_title": "Scale",
            "feat_3_body": "Bigger",
        },
        4: {"headline": "Thanks!"},
    }
    result = render_deck("midnight_tech", slot_values, str(tmp_path / "deck"))
    assert result["success"] is True
    assert len(result["html_paths"]) == 5
    # the cover HTML must contain the substituted title text
    cover_html = Path(result["html_paths"][0]).read_text(encoding="utf-8")
    assert "Hello World" in cover_html
    assert "Your Title Goes Here" not in cover_html


def test_render_deck_rejects_missing_required(tmp_path: Path) -> None:
    # nothing supplied -> every required slot is missing
    result = render_deck("midnight_tech", {}, str(tmp_path / "deck"))
    assert result["success"] is False
    assert "errors" in result
    assert any("title" in e for e in result["errors"])


@pytest.mark.asyncio
async def test_create_deck_end_to_end(tmp_path: Path) -> None:
    pytest.importorskip("playwright.async_api")
    slot_values = {
        0: {"title": "Smoke Test"},
        1: {"heading": "Agenda", "item_1": "One", "item_2": "Two", "item_3": "Three"},
        2: {"title": "Section"},
        3: {
            "heading": "Three features",
            "feat_1_title": "A",
            "feat_1_body": "alpha",
            "feat_2_title": "B",
            "feat_2_body": "beta",
            "feat_3_title": "C",
            "feat_3_body": "gamma",
        },
        4: {"headline": "Done"},
    }
    out_pptx = tmp_path / "deck.pptx"
    result = await create_deck_from_template(
        "midnight_tech", slot_values, str(out_pptx)
    )
    assert result["success"] is True
    assert result["pptx"]["success"] is True
    assert out_pptx.exists()
    pres = Presentation(out_pptx)
    assert len(pres.slides) == 5
    # the substituted title must appear as a text shape on slide 0
    texts = [s.text_frame.text for s in pres.slides[0].shapes if s.has_text_frame]
    assert any("Smoke Test" in t for t in texts)


@pytest.mark.asyncio
async def test_export_only_wrapper(tmp_path: Path) -> None:
    pytest.importorskip("playwright.async_api")
    fixture = (
        Path(__file__).parent.parent.parent.parent / "resources" / "html_to_pptx"
    )
    pages = sorted(fixture.glob("page_*.html"))
    out = tmp_path / "wrapper.pptx"
    result = await export_deck_to_pptx([str(p) for p in pages], str(out))
    assert result["success"] is True
    assert out.exists() and out.stat().st_size > 5_000
