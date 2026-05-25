"""Smoke test for the html_to_pptx converter."""

from __future__ import annotations

import asyncio
from pathlib import Path

import pytest
from pptx import Presentation

from xagent.core.tools.core.html_to_pptx import (
    HtmlToPptxConverter,
    _dedupe_parent_child_text,
    _is_transparent,
    _parse_color,
    _parse_linear_gradient,
)

FIXTURE_DIR = Path(__file__).parent.parent.parent.parent / "resources" / "html_to_pptx"


def test_parse_color_rgb_and_hex() -> None:
    c = _parse_color("rgb(10, 20, 30)")
    assert c is not None and (c[0], c[1], c[2]) == (10, 20, 30)
    c = _parse_color("#0a141e")
    assert c is not None and (c[0], c[1], c[2]) == (10, 20, 30)
    assert _parse_color(None) is None
    assert _parse_color("not-a-color") is None


def test_is_transparent() -> None:
    assert _is_transparent("rgba(0, 0, 0, 0)")
    assert _is_transparent("transparent")
    assert _is_transparent(None)
    assert not _is_transparent("rgb(10, 20, 30)")
    assert _is_transparent("rgba(255, 0, 0, 0)")  # alpha == 0


def test_parse_linear_gradient() -> None:
    stops = _parse_linear_gradient("linear-gradient(135deg, rgb(15, 23, 42), rgb(30, 58, 138))")
    assert len(stops) == 2
    assert (stops[0][0], stops[0][1], stops[0][2]) == (15, 23, 42)
    assert (stops[1][0], stops[1][1], stops[1][2]) == (30, 58, 138)
    # rejects non-gradient strings
    assert _parse_linear_gradient("none") == []


def test_dedupe_parent_child_text() -> None:
    parent = {"text": "Hello", "x": 0, "y": 0, "w": 200, "h": 80}
    child = {"text": "Hello", "x": 10, "y": 10, "w": 100, "h": 30}
    cleaned = _dedupe_parent_child_text([parent, child])
    # parent loses its text; child keeps it
    assert cleaned[0]["text"] == ""
    assert cleaned[1]["text"] == "Hello"


@pytest.mark.asyncio
async def test_convert_end_to_end(tmp_path: Path) -> None:
    pytest.importorskip("playwright.async_api")
    pages = [FIXTURE_DIR / "page_0.html", FIXTURE_DIR / "page_1.html"]
    out = tmp_path / "out.pptx"
    result = await HtmlToPptxConverter().convert(pages, out)
    assert result["success"] is True
    assert result["slide_count"] == 2
    assert out.exists() and out.stat().st_size > 5_000

    pres = Presentation(out)
    assert len(pres.slides) == 2
    # First slide must contain the title text we put in the fixture.
    texts = [
        shape.text_frame.text
        for shape in pres.slides[0].shapes
        if shape.has_text_frame
    ]
    assert any("Test Deck" in t for t in texts)
    # Each card body should produce a separate text box.
    assert any("First card body text" in t for t in texts)


if __name__ == "__main__":  # quick manual run
    asyncio.run(
        HtmlToPptxConverter().convert(
            [FIXTURE_DIR / "page_0.html", FIXTURE_DIR / "page_1.html"],
            Path("/tmp/html_to_pptx_smoke.pptx"),
        )
    )
